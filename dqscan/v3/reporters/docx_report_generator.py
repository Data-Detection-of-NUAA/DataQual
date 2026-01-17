# -*- coding: utf-8 -*-

"""
Word（.docx）报告生成器。

依赖：python-docx。
如果依赖缺失，上层 `DetectionReportGenerator.generate_full_report` 会捕获异常并在 paths 中返回 `docx_report_error`。
"""

from __future__ import annotations

import os
from datetime import datetime
from typing import Any

from .base_reporter import BaseReporter
from .scoring import AdversarialScorer, DirtyDataScorer, DistributionScorer, PhysicsScorer

try:
    from docx import Document  # type: ignore
    from docx.enum.table import WD_TABLE_ALIGNMENT  # type: ignore
    from docx.enum.text import WD_ALIGN_PARAGRAPH  # type: ignore
    from docx.shared import Pt  # type: ignore

    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False


class DocxReportGenerator(BaseReporter):
    SCORERS = {
        "distribution": DistributionScorer,
        "dirty_data": DirtyDataScorer,
        "adversarial": AdversarialScorer,
        "physics": PhysicsScorer,
    }

    def __init__(self, output_dir: str):
        if not DOCX_AVAILABLE:
            raise ImportError("python-docx未安装。请运行: pip install python-docx")
        super().__init__(output_dir)
        self.doc = Document()
        self._setup_styles()

    def _setup_styles(self) -> None:
        style = self.doc.styles["Normal"]
        style.font.name = "Calibri"
        style.font.size = Pt(10.5)

    def _get_scorer(self, module_type: str):
        cls = self.SCORERS.get(module_type)
        return cls() if cls else None

    def generate_report(self, results: dict[str, Any], report_name: str, module_type: str) -> str:
        scorer = self._get_scorer(module_type)
        scoring = scorer.calculate_total_score(results) if scorer else {"total_score": 0, "grade": "N/A", "data_type_scores": {}}
        score = float(scoring.get("total_score", 0) or 0)
        grade = self.get_grade_full(score) if scorer else "N/A"

        self._add_cover(module_type, score, grade)
        self._add_overview(results, scoring, module_type)
        self._add_details(results, module_type)
        self._add_scoring_explain(module_type, scorer)

        path = os.path.join(self.report_dir, f"{report_name}.docx")
        self.doc.save(path)
        return path

    def _add_cover(self, module_type: str, score: float, grade: str) -> None:
        module_names = {
            "distribution": "分布偏差检测",
            "dirty_data": "脏数据扫描",
            "adversarial": "对抗性检测",
            "physics": "物理保真度检测",
            "all": "综合质量检测",
        }
        title = self.doc.add_heading("数据质量检测报告", level=0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sub = self.doc.add_paragraph(f"—— {module_names.get(module_type, module_type)} ——")
        sub.alignment = WD_ALIGN_PARAGRAPH.CENTER

        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(f"总分：{score:.1f}    等级：{grade}")
        r.bold = True
        r.font.size = Pt(14)

        meta = self.doc.add_paragraph(f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.doc.add_page_break()

    def _add_overview(self, results: dict[str, Any], scoring: dict[str, Any], module_type: str) -> None:
        self.doc.add_heading("1. 总览", level=1)
        table = self.doc.add_table(rows=1, cols=6)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        hdr = table.rows[0].cells
        hdr[0].text = "数据类型"
        hdr[1].text = "算法"
        hdr[2].text = "问题数"
        hdr[3].text = "问题比例"
        hdr[4].text = "评分"
        hdr[5].text = "状态"

        dt_scores = scoring.get("data_type_scores", {}) if isinstance(scoring.get("data_type_scores"), dict) else {}

        for dt, dt_result in results.items():
            if not isinstance(dt_result, dict):
                continue
            row = table.add_row().cells
            row[0].text = self.DATA_TYPE_LABELS.get(dt, dt)
            row[1].text = str(dt_result.get("algorithm", "Unknown"))
            row[2].text = str(dt_result.get("total_issues", 0))
            row[3].text = str(round(float(dt_result.get("issue_percentage", 0) or 0), 4))
            row[4].text = str(dt_scores.get(dt, {}).get("score", ""))
            row[5].text = "有风险" if dt_result.get("has_issues") else "正常"

        self.doc.add_paragraph("")

    def _add_details(self, results: dict[str, Any], module_type: str) -> None:
        self.doc.add_heading("2. 详情", level=1)
        for dt, dt_result in results.items():
            if not isinstance(dt_result, dict):
                continue
            self.doc.add_heading(f"2.1 {self.DATA_TYPE_LABELS.get(dt, dt)}", level=2)
            if "error" in dt_result:
                self.doc.add_paragraph(f"错误：{dt_result['error']}")
                continue

            if module_type == "dirty_data":
                self.doc.add_paragraph(
                    f"异常率: {dt_result.get('anomaly_rate', 0):.4f}  缺失率: {dt_result.get('missing_rate', 0):.4f}  重复率: {dt_result.get('duplicate_rate', 0):.4f}"
                )
            elif module_type == "distribution":
                self.doc.add_paragraph(
                    f"漂移: {dt_result.get('drift_detected', False)}  p_value: {dt_result.get('p_value', 1.0)}"
                )
            elif module_type == "adversarial":
                self.doc.add_paragraph(
                    f"攻击成功率: {dt_result.get('attack_success_rate', 0):.4f}  鲁棒性: {dt_result.get('robustness_score', 1.0):.4f}"
                )
            elif module_type == "physics":
                self.doc.add_paragraph(
                    f"违规率: {dt_result.get('violation_rate', 0):.4f}  因果得分: {dt_result.get('causality_score', 1.0):.4f}"
                )

            issues = dt_result.get("detailed_issues", [])
            if isinstance(issues, list) and issues:
                self.doc.add_paragraph("示例问题：")
                for item in issues[:20]:
                    self.doc.add_paragraph(f"- {item.get('issue_type')}: {item.get('data_id')}", style="List Bullet")

        self.doc.add_page_break()

    def _add_scoring_explain(self, module_type: str, scorer: Any) -> None:
        self.doc.add_heading("3. 评分说明", level=1)
        if scorer is None:
            self.doc.add_paragraph("该模块无评分器。")
            return
        exp = scorer.get_scoring_explanation()
        self.doc.add_paragraph(f"方法：{exp.get('method')}")
        self.doc.add_paragraph(f"说明：{exp.get('description')}")
