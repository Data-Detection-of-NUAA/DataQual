"""Word 报告生成器（可视化风格）。"""

from __future__ import annotations

import os
from datetime import datetime
from typing import Any

from .base_reporter import BaseReporter
from .scoring import AdversarialScorer, DirtyDataScorer, DistributionScorer, PhysicsScorer

try:
    from docx import Document  # type: ignore
    from docx.enum.table import WD_TABLE_ALIGNMENT  # type: ignore
    from docx.enum.text import WD_ALIGN_PARAGRAPH  # type: ignore
    from docx.oxml import OxmlElement  # type: ignore
    from docx.oxml.ns import qn  # type: ignore
    from docx.shared import Inches, Pt, RGBColor  # type: ignore

    DOCX_AVAILABLE = True
except Exception:
    DOCX_AVAILABLE = False
    Document = None  # type: ignore
    WD_TABLE_ALIGNMENT = None  # type: ignore
    WD_ALIGN_PARAGRAPH = None  # type: ignore
    OxmlElement = None  # type: ignore
    qn = None  # type: ignore
    Inches = None  # type: ignore
    Pt = None  # type: ignore
    RGBColor = None  # type: ignore


def _rgb(r: int, g: int, b: int):
    # 允许在 python-docx 缺失时导入本模块（上层会捕获并在报告中降级）
    if DOCX_AVAILABLE and RGBColor:
        return RGBColor(r, g, b)  # type: ignore[misc]
    return (r, g, b)

try:
    from .chart_generator import ChartGenerator, MATPLOTLIB_AVAILABLE
except Exception:
    ChartGenerator = None  # type: ignore
    MATPLOTLIB_AVAILABLE = False


class DocxReportGenerator(BaseReporter):
    """Word 文档报告生成器 - 增强可视化版本（对齐参考工程 output/reports 风格）。"""

    SCORERS = {
        "distribution": DistributionScorer,
        "dirty_data": DirtyDataScorer,
        "adversarial": AdversarialScorer,
        "physics": PhysicsScorer,
    }

    COLORS = {
        "excellent": _rgb(0x28, 0xA7, 0x45),  # 绿色
        "good": _rgb(0x17, 0xA2, 0xB8),  # 蓝色
        "warning": _rgb(0xFF, 0xC1, 0x07),  # 黄色
        "danger": _rgb(0xDC, 0x35, 0x45),  # 红色
        "neutral": _rgb(0x6C, 0x75, 0x7D),  # 灰色
        "primary": _rgb(0x1F, 0x49, 0x7D),  # 深蓝
    }

    def __init__(self, output_dir: str):
        if not DOCX_AVAILABLE:
            raise ImportError("python-docx未安装。请运行: pip install python-docx")
        super().__init__(output_dir)
        self.doc = Document()
        self._setup_styles()
        self.chart_gen = None
        if MATPLOTLIB_AVAILABLE and ChartGenerator:
            self.chart_gen = ChartGenerator(os.path.join(output_dir, "charts"))

    def _setup_styles(self) -> None:
        style = self.doc.styles["Normal"]
        style.font.name = "Calibri"
        style.font.size = Pt(10.5)
        try:
            style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        except Exception:
            pass

        for i, size in [(1, 16), (2, 14), (3, 12)]:
            try:
                h = self.doc.styles[f"Heading {i}"]
                h.font.name = "Calibri"
                h.font.size = Pt(size)
                h.font.bold = True
                h.font.color.rgb = self.COLORS["primary"]
                h._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
            except Exception:
                continue

    def generate_report(self, results: dict[str, Any], report_name: str, module_type: str) -> str:
        scorer = self._get_scorer(module_type)
        scoring = (
            scorer.calculate_total_score(results)
            if scorer
            else {"total_score": 0, "grade": "N/A", "data_type_scores": {}}
        )
        score = float(scoring.get("total_score", 0) or 0)
        grade = self.get_grade_full(score) if scorer else "N/A"

        self._generate_visual_cover(score, grade, module_type, results, scoring)
        self._generate_visual_overview(results, scoring, module_type)

        if module_type == "distribution":
            self._generate_distribution_section(results, scoring)
        elif module_type == "dirty_data":
            self._generate_dirty_data_section(results, scoring)
        elif module_type == "adversarial":
            self._generate_adversarial_section(results, scoring)
        elif module_type == "physics":
            self._generate_physics_section(results, scoring)

        self._generate_scoring_explanation(module_type)

        report_path = os.path.join(self.report_dir, f"{report_name}.docx")
        self.doc.save(report_path)

        if self.chart_gen:
            self.chart_gen.cleanup()
        return report_path

    def _get_scorer(self, module_type: str):
        scorer_class = self.SCORERS.get(module_type)
        return scorer_class() if scorer_class else None

    def _generate_visual_cover(
        self,
        score: float,
        grade: str,
        module_type: str,
        results: dict[str, Any],
        scoring: dict[str, Any],
    ) -> None:
        module_names = {
            "distribution": "分布偏差检测",
            "dirty_data": "脏数据扫描",
            "adversarial": "对抗性检测",
            "physics": "物理保真度检测",
            "all": "综合质量检测",
        }

        self._add_paragraph("")
        title = self.doc.add_heading("数据质量检测报告", level=0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        sub = self.doc.add_paragraph()
        sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = sub.add_run(f"—— {module_names.get(module_type, '质量检测')} ——")
        run.font.size = Pt(14)
        run.font.color.rgb = self.COLORS["neutral"]

        self._add_paragraph("")
        self._add_metric_cards(score, grade, results)
        self._add_paragraph("")

        if self.chart_gen and scoring.get("data_type_scores"):
            try:
                dt_scores = {
                    str(k): float(v.get("score", 0) or 0)
                    for k, v in (scoring.get("data_type_scores") or {}).items()
                    if isinstance(v, dict)
                }
                chart_path = self.chart_gen.generate_data_type_bars(dt_scores, "各数据类型评分")
                if chart_path and os.path.exists(chart_path):
                    p = self.doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run = p.add_run()
                    run.add_picture(chart_path, width=Inches(5))
            except Exception:
                pass

        self._add_paragraph("")
        meta_table = self.doc.add_table(rows=1, cols=2)
        meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cells = meta_table.rows[0].cells
        cells[0].text = "报告来源: dqscan"
        cells[1].text = f"生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
        for cell in cells:
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in cell.paragraphs[0].runs:
                r.font.size = Pt(9)
                r.font.color.rgb = self.COLORS["neutral"]

        self.doc.add_page_break()

    def _add_metric_cards(self, score: float, grade: str, results: dict[str, Any]) -> None:
        total_issues = sum(
            int(v.get("total_issues", 0) or 0) for v in results.values() if isinstance(v, dict)
        )
        data_type_count = len([k for k, v in results.items() if isinstance(v, dict)])

        table = self.doc.add_table(rows=2, cols=4)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        cards = [
            ("总分", f"{score:.0f}", self._get_score_color(score)),
            ("等级", grade, self._get_grade_color(grade)),
            ("数据类型", str(data_type_count), self.COLORS["primary"]),
            ("问题数", str(total_issues), self.COLORS["danger"] if total_issues > 0 else self.COLORS["excellent"]),
        ]

        for idx, (label, value, color) in enumerate(cards):
            value_cell = table.rows[0].cells[idx]
            value_p = value_cell.paragraphs[0]
            value_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            value_run = value_p.add_run(value)
            value_run.font.size = Pt(28)
            value_run.font.bold = True
            value_run.font.color.rgb = color

            label_cell = table.rows[1].cells[idx]
            label_p = label_cell.paragraphs[0]
            label_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            label_run = label_p.add_run(label)
            label_run.font.size = Pt(10)
            label_run.font.color.rgb = self.COLORS["neutral"]

    def _generate_visual_overview(self, results: dict[str, Any], scoring: dict[str, Any], module_type: str) -> None:
        self.doc.add_heading("1. 检测总览", level=1)
        if scoring.get("data_type_scores"):
            self._add_data_type_summary_table(results, scoring, module_type)
        if self.chart_gen and module_type == "dirty_data":
            self._add_rate_comparison_chart(results)
        self._add_paragraph("")

    def _add_data_type_summary_table(self, results: dict[str, Any], scoring: dict[str, Any], module_type: str) -> None:
        labels = {"tabular": "表格数据", "timeseries": "时序数据", "image": "图像数据", "text": "文本数据"}

        if module_type == "distribution":
            headers = ["数据类型", "算法", "p值", "漂移检测", "评分", "状态"]
        elif module_type == "adversarial":
            headers = ["数据类型", "算法", "攻击成功率", "鲁棒性", "评分", "状态"]
        elif module_type == "physics":
            headers = ["数据类型", "算法", "违规率", "保真度", "评分", "状态"]
        else:
            headers = ["数据类型", "算法", "异常率", "缺失率", "重复率", "评分", "状态"]

        table = self.doc.add_table(rows=1, cols=len(headers))
        table.style = "Table Grid"

        for i, h in enumerate(headers):
            cell = table.rows[0].cells[i]
            cell.text = h
            self._set_cell_shading(cell, "E9ECEF")
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.bold = True
                    r.font.size = Pt(9)

        data_type_scores = scoring.get("data_type_scores") if isinstance(scoring.get("data_type_scores"), dict) else {}
        for dt, dt_results in results.items():
            if not isinstance(dt_results, dict):
                continue

            dt_score_info = data_type_scores.get(dt, {}) if isinstance(data_type_scores, dict) else {}
            score = float(dt_score_info.get("score", 0) or 0)

            row = table.add_row()
            cells = row.cells

            if module_type == "distribution":
                p_value = float(dt_score_info.get("p_value", 1.0) or 1.0)
                drift = bool(dt_results.get("drift_detected", False))
                values = [
                    labels.get(dt, dt),
                    str(dt_results.get("algorithm", "N/A"))[:20],
                    f"{p_value:.4f}",
                    "是" if drift else "否",
                    f"{score:.0f}",
                    self._get_status_icon(score),
                ]
            elif module_type == "adversarial":
                attack_rate = float(dt_results.get("attack_success_rate", 0) or 0)
                robustness = float(dt_results.get("robustness_score", 1 - attack_rate) or 0)
                values = [
                    labels.get(dt, dt),
                    str(dt_results.get("algorithm", "N/A"))[:20],
                    f"{attack_rate:.1%}",
                    f"{robustness:.1%}",
                    f"{score:.0f}",
                    self._get_status_icon(score),
                ]
            elif module_type == "physics":
                violation_rate = float(dt_results.get("violation_rate", 0) or 0)
                fidelity_level = str(dt_score_info.get("fidelity_level", "N/A"))
                values = [
                    labels.get(dt, dt),
                    str(dt_results.get("algorithm", "N/A"))[:20],
                    f"{violation_rate:.1%}",
                    fidelity_level,
                    f"{score:.0f}",
                    self._get_status_icon(score),
                ]
            else:
                values = [
                    labels.get(dt, dt),
                    str(dt_results.get("algorithm", "N/A"))[:15],
                    f"{float(dt_results.get('anomaly_rate', 0) or 0):.1%}",
                    f"{float(dt_results.get('missing_rate', 0) or 0):.1%}",
                    f"{float(dt_results.get('duplicate_rate', 0) or 0):.1%}",
                    f"{score:.0f}",
                    self._get_status_icon(score),
                ]

            for i, val in enumerate(values):
                cells[i].text = str(val)
                for p in cells[i].paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs:
                        r.font.size = Pt(9)
                        if i == len(values) - 1:
                            r.font.color.rgb = self._get_score_color(score)
                        elif i == len(values) - 2:
                            r.font.bold = True
                            r.font.color.rgb = self._get_score_color(score)

        self._add_paragraph("")

    def _add_rate_comparison_chart(self, results: dict[str, Any]) -> None:
        try:
            data_type_rates: dict[str, dict[str, float]] = {}
            for dt, dt_results in results.items():
                if isinstance(dt_results, dict):
                    data_type_rates[dt] = {
                        "anomaly_rate": float(dt_results.get("anomaly_rate", 0) or 0),
                        "missing_rate": float(dt_results.get("missing_rate", 0) or 0),
                        "duplicate_rate": float(dt_results.get("duplicate_rate", 0) or 0),
                    }

            if self.chart_gen and data_type_rates:
                chart_path = self.chart_gen.generate_rate_comparison(data_type_rates)
                if chart_path and os.path.exists(chart_path):
                    p = self.doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run = p.add_run()
                    run.add_picture(chart_path, width=Inches(5.5))
        except Exception:
            pass

    def _generate_distribution_section(self, results: dict[str, Any], scoring: dict[str, Any]) -> None:
        self.doc.add_heading("2. 检测详情", level=1)
        labels = {"tabular": "表格数据", "timeseries": "时序数据", "image": "图像数据", "text": "文本数据"}

        data_type_scores = scoring.get("data_type_scores") if isinstance(scoring.get("data_type_scores"), dict) else {}
        for idx, (dt, dt_results) in enumerate(results.items(), 1):
            if not isinstance(dt_results, dict):
                continue
            dt_score = data_type_scores.get(dt, {}) if isinstance(data_type_scores, dict) else {}
            self.doc.add_heading(f"2.{idx} {labels.get(dt, dt)}", level=2)
            self._add_distribution_metrics_table(dt_results, dt_score)
            self._add_issue_summary(dt_results)
            self._add_paragraph("")

        self.doc.add_page_break()

    def _add_distribution_metrics_table(self, dt_results: dict[str, Any], dt_score: dict[str, Any]) -> None:
        table = self.doc.add_table(rows=2, cols=5)
        table.style = "Table Grid"

        p_value = float(dt_score.get("p_value", 1.0) or 1.0)
        drift_detected = bool(dt_results.get("drift_detected", False))
        significance = str(dt_score.get("significance", "N/A"))
        score = float(dt_score.get("score", 0) or 0)

        info_data = [
            ["算法", str(dt_results.get("algorithm", "N/A"))[:25]],
            ["p值", f"{p_value:.6f}"],
            ["显著性", significance],
            ["漂移检测", "是" if drift_detected else "否"],
            ["评分", f"{score:.0f}分"],
        ]

        for i, (label, value) in enumerate(info_data):
            table.rows[0].cells[i].text = label
            table.rows[1].cells[i].text = value
            for row in table.rows:
                for p in row.cells[i].paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs:
                        r.font.size = Pt(9)
            self._set_cell_shading(table.rows[0].cells[i], "F0F0F0")

        drift_cell = table.rows[1].cells[3]
        for p in drift_cell.paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.color.rgb = self.COLORS["danger"] if drift_detected else self.COLORS["excellent"]

        self._add_paragraph("")

    def _generate_dirty_data_section(self, results: dict[str, Any], scoring: dict[str, Any]) -> None:
        self.doc.add_heading("2. 检测详情", level=1)
        labels = {"tabular": "表格数据", "timeseries": "时序数据", "image": "图像数据", "text": "文本数据"}

        data_type_scores = scoring.get("data_type_scores") if isinstance(scoring.get("data_type_scores"), dict) else {}
        for idx, (dt, dt_results) in enumerate(results.items(), 1):
            if not isinstance(dt_results, dict):
                continue
            dt_score = data_type_scores.get(dt, {}) if isinstance(data_type_scores, dict) else {}
            self.doc.add_heading(f"2.{idx} {labels.get(dt, dt)}", level=2)
            self._add_rate_progress_table(dt_results)
            self._add_issue_summary(dt_results)
            self._add_paragraph("")

        self.doc.add_page_break()

    def _generate_adversarial_section(self, results: dict[str, Any], scoring: dict[str, Any]) -> None:
        self.doc.add_heading("2. 检测详情", level=1)
        labels = {"tabular": "表格数据", "timeseries": "时序数据", "image": "图像数据", "text": "文本数据"}

        data_type_scores = scoring.get("data_type_scores") if isinstance(scoring.get("data_type_scores"), dict) else {}
        for idx, (dt, dt_results) in enumerate(results.items(), 1):
            if not isinstance(dt_results, dict):
                continue
            dt_score = data_type_scores.get(dt, {}) if isinstance(data_type_scores, dict) else {}
            self.doc.add_heading(f"2.{idx} {labels.get(dt, dt)}", level=2)
            self._add_adversarial_metrics_table(dt_results, dt_score)
            self._add_issue_summary(dt_results)
            self._add_paragraph("")

        self.doc.add_page_break()

    def _generate_physics_section(self, results: dict[str, Any], scoring: dict[str, Any]) -> None:
        self.doc.add_heading("2. 检测详情", level=1)
        labels = {"tabular": "表格数据", "timeseries": "时序数据", "image": "图像数据", "text": "文本数据"}

        data_type_scores = scoring.get("data_type_scores") if isinstance(scoring.get("data_type_scores"), dict) else {}
        for idx, (dt, dt_results) in enumerate(results.items(), 1):
            if not isinstance(dt_results, dict):
                continue
            dt_score = data_type_scores.get(dt, {}) if isinstance(data_type_scores, dict) else {}
            self.doc.add_heading(f"2.{idx} {labels.get(dt, dt)}", level=2)
            self._add_physics_metrics_table(dt_results, dt_score)
            self._add_constraints_checked(dt_results)
            self._add_issue_summary(dt_results)
            self._add_paragraph("")

        self.doc.add_page_break()

    def _add_physics_metrics_table(self, dt_results: dict[str, Any], dt_score: dict[str, Any]) -> None:
        table = self.doc.add_table(rows=2, cols=5)
        table.style = "Table Grid"

        violation_rate = float(dt_results.get("violation_rate", 0) or 0)
        causality_score = float(dt_results.get("causality_score", 1.0) or 0)
        fidelity_level = str(dt_score.get("fidelity_level", "N/A"))
        score = float(dt_score.get("score", 0) or 0)

        info_data = [
            ["算法", str(dt_results.get("algorithm", "N/A"))[:25]],
            ["违规率", f"{violation_rate:.1%}"],
            ["因果得分", f"{causality_score:.1%}"],
            ["保真度", fidelity_level],
            ["评分", f"{score:.0f}分"],
        ]

        for i, (label, value) in enumerate(info_data):
            table.rows[0].cells[i].text = label
            table.rows[1].cells[i].text = value
            for row in table.rows:
                for p in row.cells[i].paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs:
                        r.font.size = Pt(9)
            self._set_cell_shading(table.rows[0].cells[i], "F8F9FA")

        fidelity_cell = table.rows[1].cells[3]
        for p in fidelity_cell.paragraphs:
            for r in p.runs:
                r.font.bold = True
                if fidelity_level == "高保真":
                    r.font.color.rgb = self.COLORS["excellent"]
                elif fidelity_level == "中等保真":
                    r.font.color.rgb = self.COLORS["good"]
                elif fidelity_level == "低保真":
                    r.font.color.rgb = self.COLORS["warning"]
                else:
                    r.font.color.rgb = self.COLORS["danger"]

        self._add_paragraph("")

    def _add_constraints_checked(self, dt_results: dict[str, Any]) -> None:
        constraints = dt_results.get("constraints_checked", [])
        violation_types = dt_results.get("violation_types", {}) if isinstance(dt_results.get("violation_types"), dict) else {}
        if not constraints:
            return

        table = self.doc.add_table(rows=1, cols=3)
        table.style = "Table Grid"

        headers = ["约束项", "违规数", "状态"]
        for i, h in enumerate(headers):
            cell = table.rows[0].cells[i]
            cell.text = h
            self._set_cell_shading(cell, "F0F0F0")
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.bold = True
                    r.font.size = Pt(9)

        for constraint in constraints:
            violations = int(violation_types.get(constraint, 0) or 0)
            status = "通过" if violations == 0 else "违规"
            row = table.add_row()
            row.cells[0].text = str(constraint)
            row.cells[1].text = str(violations)
            row.cells[2].text = status

            for j, cell in enumerate(row.cells):
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs:
                        r.font.size = Pt(9)
                        if j == 2:
                            r.font.color.rgb = self.COLORS["excellent"] if violations == 0 else self.COLORS["danger"]

        self._add_paragraph("")

    def _add_adversarial_metrics_table(self, dt_results: dict[str, Any], dt_score: dict[str, Any]) -> None:
        table = self.doc.add_table(rows=2, cols=5)
        table.style = "Table Grid"

        attack_rate = float(dt_results.get("attack_success_rate", 0) or 0)
        robustness = float(dt_results.get("robustness_score", 1 - attack_rate) or 0)
        total_samples = int(dt_results.get("total_samples", 0) or 0)
        total_issues = int(dt_results.get("total_issues", 0) or 0)
        score = float(dt_score.get("score", 0) or 0)

        info_data = [
            ["算法", str(dt_results.get("algorithm", "N/A"))[:25]],
            ["攻击成功率", f"{attack_rate:.1%}"],
            ["鲁棒性", f"{robustness:.1%}"],
            ["成功攻击数", f"{total_issues}/{total_samples}"],
            ["评分", f"{score:.0f}分"],
        ]

        for i, (label, value) in enumerate(info_data):
            table.rows[0].cells[i].text = label
            table.rows[1].cells[i].text = value
            for row in table.rows:
                for p in row.cells[i].paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs:
                        r.font.size = Pt(9)
            self._set_cell_shading(table.rows[0].cells[i], "F8F9FA")

        robustness_cell = table.rows[1].cells[2]
        for p in robustness_cell.paragraphs:
            for r in p.runs:
                r.font.bold = True
                if robustness >= 0.9:
                    r.font.color.rgb = self.COLORS["excellent"]
                elif robustness >= 0.7:
                    r.font.color.rgb = self.COLORS["good"]
                elif robustness >= 0.5:
                    r.font.color.rgb = self.COLORS["warning"]
                else:
                    r.font.color.rgb = self.COLORS["danger"]

        self._add_paragraph("")

    def _add_rate_progress_table(self, dt_results: dict[str, Any]) -> None:
        rates = [
            ("异常率", float(dt_results.get("anomaly_rate", 0) or 0), 0.05, 0.10),
            ("缺失率", float(dt_results.get("missing_rate", 0) or 0), 0.01, 0.05),
            ("重复率", float(dt_results.get("duplicate_rate", 0) or 0), 0.05, 0.15),
        ]
        if "label_mismatch_rate" in dt_results:
            rates.append(("疑似错标率", float(dt_results.get("label_mismatch_rate", 0) or 0), 0.02, 0.05))

        table = self.doc.add_table(rows=len(rates) + 1, cols=4)
        table.style = "Table Grid"

        headers = ["指标", "数值", "阈值", "状态"]
        for i, h in enumerate(headers):
            table.rows[0].cells[i].text = h
            self._set_cell_shading(table.rows[0].cells[i], "F0F0F0")
            for p in table.rows[0].cells[i].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.bold = True
                    r.font.size = Pt(9)

        for i, (name, rate, warn_thresh, danger_thresh) in enumerate(rates, 1):
            if rate <= warn_thresh:
                status = "正常"
                color = self.COLORS["excellent"]
            elif rate <= danger_thresh:
                status = "警告"
                color = self.COLORS["warning"]
            else:
                status = "严重"
                color = self.COLORS["danger"]

            row = table.rows[i]
            row.cells[0].text = name
            row.cells[1].text = f"{rate:.2%}"
            row.cells[2].text = f"<{warn_thresh:.0%}正常"
            row.cells[3].text = status

            for j, cell in enumerate(row.cells):
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs:
                        r.font.size = Pt(9)
                        if j == 3:
                            r.font.color.rgb = color

        self._add_paragraph("")

    def _add_issue_summary(self, dt_results: dict[str, Any]) -> None:
        issues = dt_results.get("detailed_issues", [])
        if not isinstance(issues, list) or not issues:
            return

        issue_counts: dict[str, int] = {}
        for issue in issues:
            if not isinstance(issue, dict):
                continue
            issue_type = str(issue.get("issue_type", "未知"))
            issue_counts[issue_type] = issue_counts.get(issue_type, 0) + 1

        summary_parts = [f"{k}({v})" for k, v in issue_counts.items()]
        self._add_paragraph(f"检测到 {len(issues)} 个问题: " + ", ".join(summary_parts), bold=False)

    def _generate_scoring_explanation(self, module_type: str) -> None:
        self.doc.add_heading("3. 评分说明", level=1)

        if module_type == "distribution":
            self._add_distribution_scoring_table()
        elif module_type == "dirty_data":
            self._add_dirty_data_scoring_table()
        elif module_type == "adversarial":
            self._add_adversarial_scoring_table()
        elif module_type == "physics":
            self._add_physics_scoring_table()

        self.doc.add_heading("3.2 评级标准", level=2)
        grade_table = self.doc.add_table(rows=2, cols=5)
        grade_table.style = "Table Grid"

        grades = [("S", "90-100"), ("A", "80-89"), ("B", "70-79"), ("C", "60-69"), ("D", "0-59")]
        for i, (g, r) in enumerate(grades):
            grade_table.rows[0].cells[i].text = f"{g}级"
            grade_table.rows[1].cells[i].text = f"{r}分"
            for row in grade_table.rows:
                for p in row.cells[i].paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in p.runs:
                        run.font.size = Pt(9)

    def _add_distribution_scoring_table(self) -> None:
        self.doc.add_heading("3.1 分布偏差评分规则", level=2)
        self._add_paragraph("评分基于统计显著性p值:")

        table = self.doc.add_table(rows=5, cols=3)
        table.style = "Table Grid"
        data = [
            ["p值范围", "显著性", "评分"],
            ["p >= 0.05", "不显著", "100分"],
            ["0.01 <= p < 0.05", "弱显著", "80分"],
            ["0.001 <= p < 0.01", "显著", "60分"],
            ["p < 0.001", "高度显著", "40分"],
        ]
        self._fill_table_grid(table, data, header_row=0)
        self._add_paragraph("")

    def _add_dirty_data_scoring_table(self) -> None:
        self.doc.add_heading("3.1 脏数据评分规则", level=2)
        self._add_paragraph("评分采用扣分制（满分100分）:")

        table = self.doc.add_table(rows=4, cols=5)
        table.style = "Table Grid"
        data = [
            ["指标", "<=阈值1", "阈值1~2", "阈值2~3", ">阈值3"],
            ["异常率", "<=5%: 0", "5-10%: -20", "10-20%: -40", ">20%: -60"],
            ["缺失率", "<=1%: 0", "1-5%: -10", "5-15%: -25", ">15%: -40"],
            ["重复率", "<=5%: 0", "5-15%: -10", "15-30%: -20", ">30%: -30"],
        ]
        self._fill_table_grid(table, data, header_row=0)
        self._add_paragraph("")

    def _add_adversarial_scoring_table(self) -> None:
        self.doc.add_heading("3.1 对抗性检测评分规则", level=2)
        self._add_paragraph("评分基于攻击成功率:")

        table = self.doc.add_table(rows=5, cols=3)
        table.style = "Table Grid"
        data = [
            ["攻击成功率", "鲁棒性等级", "评分"],
            ["<= 10%", "高鲁棒", "100分"],
            ["10% ~ 30%", "中等鲁棒", "80分"],
            ["30% ~ 50%", "低鲁棒", "60分"],
            ["> 50%", "极低鲁棒", "40分"],
        ]
        self._fill_table_grid(table, data, header_row=0)
        self._add_paragraph("")

    def _add_physics_scoring_table(self) -> None:
        self.doc.add_heading("3.1 物理保真度评分规则", level=2)
        self._add_paragraph("评分基于约束违规率:")

        table = self.doc.add_table(rows=5, cols=3)
        table.style = "Table Grid"
        data = [
            ["违规率范围", "保真度等级", "评分"],
            ["<= 1%", "高保真", "100分"],
            ["1% ~ 5%", "中等保真", "80分"],
            ["5% ~ 15%", "低保真", "60分"],
            ["> 15%", "极低保真", "40分"],
        ]
        self._fill_table_grid(table, data, header_row=0)
        self._add_paragraph("")

    def _fill_table_grid(self, table: Any, data: list[list[str]], header_row: int = 0) -> None:
        for i, row_data in enumerate(data):
            for j, val in enumerate(row_data):
                table.rows[i].cells[j].text = str(val)
                for p in table.rows[i].cells[j].paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs:
                        r.font.size = Pt(9)
                        if i == header_row:
                            r.bold = True
            if i == header_row:
                for cell in table.rows[i].cells:
                    self._set_cell_shading(cell, "E9ECEF")

    def _add_paragraph(self, text: str, bold: bool = False):
        p = self.doc.add_paragraph(text)
        if bold:
            for r in p.runs:
                r.bold = True
        return p

    def _set_cell_shading(self, cell: Any, color: str) -> None:
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), color)
        cell._tc.get_or_add_tcPr().append(shading)

    def _get_score_color(self, score: float) -> RGBColor:
        if score >= 90:
            return self.COLORS["excellent"]
        if score >= 80:
            return self.COLORS["good"]
        if score >= 70:
            return self.COLORS["warning"]
        return self.COLORS["danger"]

    def _get_grade_color(self, grade: str) -> RGBColor:
        grade_colors = {
            "S": self.COLORS["excellent"],
            "A": self.COLORS["good"],
            "B": self.COLORS["warning"],
            "C": self.COLORS["danger"],
            "D": self.COLORS["danger"],
        }
        return grade_colors.get(grade, self.COLORS["neutral"])

    def _get_status_icon(self, score: float) -> str:
        if score >= 90:
            return "[OK]"
        if score >= 70:
            return "[!]"
        return "[X]"
