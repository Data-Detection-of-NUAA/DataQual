"""
审计报告生成器
生成Excel格式的审计报告
"""
import os
from datetime import datetime
from typing import Dict, Any, List
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from app.config.path_conf import UPLOAD_DIR


class ReportGenerator:
    """报告生成器类"""

    @staticmethod
    async def generate_report(
        task_id: int,
        audit_result: Dict[str, Any],
        rules: List,
        gdpr_detection_results: List = None,
        gdpr_report: Dict[str, Any] = None
    ) -> str:
        """
        生成Excel格式的审计报告

        Args:
            task_id: 任务ID
            audit_result: 审计结果
            rules: 使用的规则列表
            gdpr_detection_results: GDPR检测结果列表
            gdpr_report: GDPR检测报告

        Returns:
            报告文件路径
        """
        # 创建工作簿
        wb = openpyxl.Workbook()

        # 只创建GDPR检测表
        ws_gdpr = wb.active
        ws_gdpr.title = "GDPR检测"
        if gdpr_detection_results:
            ReportGenerator._create_gdpr_sheet(ws_gdpr, gdpr_detection_results, gdpr_report, task_id)
        else:
            # 如果没有检测结果，显示提示信息
            ws_gdpr['A1'] = 'GDPR检测'
            ws_gdpr['A1'].font = Font(name='微软雅黑', size=14, bold=True)
            ws_gdpr['A3'] = '未检测到GDPR敏感数据'
            ws_gdpr['A3'].font = Font(name='微软雅黑', size=11, italic=True)

        # 保存文件
        today = datetime.now()
        dir_path = os.path.join(
            UPLOAD_DIR,
            "audit",
            "reports",
            str(today.year),
            f"{today.month:02d}"
        )
        os.makedirs(dir_path, exist_ok=True)

        file_name = f"audit_report_{task_id}_{today.strftime('%Y%m%d_%H%M%S')}.xlsx"
        file_path = os.path.join(dir_path, file_name)

        wb.save(file_path)
        return file_path

    @staticmethod
    def _create_summary_sheet(ws, task_id: int, audit_result: Dict[str, Any], rules: List, gdpr_report: Dict[str, Any] = None):
        """创建概览表"""
        # 标题样式
        title_font = Font(name='微软雅黑', size=16, bold=True)
        header_font = Font(name='微软雅黑', size=12, bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
        header_alignment = Alignment(horizontal='center', vertical='center')

        # 添加标题
        ws['A1'] = '数据集合规审计报告'
        ws['A1'].font = title_font
        ws.merge_cells('A1:D1')
        ws['A1'].alignment = Alignment(horizontal='center', vertical='center')

        # 添加任务信息
        row = 3
        info_items = [
            ('任务ID:', task_id),
            ('审计时间:', datetime.now().strftime('%Y-%m-%d %H:%M:%S')),
            ('总记录数:', audit_result['total_records']),
            ('错误记录数:', audit_result['error_records']),
        ]

        for label, value in info_items:
            ws[f'A{row}'] = label
            ws[f'A{row}'].font = Font(name='微软雅黑', bold=True)
            ws[f'B{row}'] = value
            row += 1

        # 计算错误率
        error_rate = (audit_result['error_records'] / audit_result['total_records'] * 100) if audit_result['total_records'] > 0 else 0
        ws[f'A{row}'] = '错误率:'
        ws[f'A{row}'].font = Font(name='微软雅黑', bold=True)
        ws[f'B{row}'] = f"{error_rate:.2f}%"
        row += 1

        # 添加GDPR合规分数（新增）
        if gdpr_report:
            ws[f'A{row}'] = 'GDPR合规分数:'
            ws[f'A{row}'].font = Font(name='微软雅黑', bold=True)
            score = gdpr_report.get('compliance_score', 0)
            ws[f'B{row}'] = f"{score:.1f}/100"
            # 根据分数设置颜色
            if score >= 80:
                ws[f'B{row}'].font = Font(name='微软雅黑', bold=True, color='00B050')  # 绿色
            elif score >= 60:
                ws[f'B{row}'].font = Font(name='微软雅黑', bold=True, color='FFC000')  # 橙色
            else:
                ws[f'B{row}'].font = Font(name='微软雅黑', bold=True, color='FF0000')  # 红色
            row += 1

        # 添加使用的规则
        row += 1
        ws[f'A{row}'] = '使用的审计规则'
        ws[f'A{row}'].font = header_font
        ws[f'A{row}'].fill = header_fill
        ws[f'A{row}'].alignment = header_alignment
        ws.merge_cells(f'A{row}:D{row}')

        row += 1
        headers = ['规则编码', '规则名称', '规则类型', '严重级别']
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=row, column=col)
            cell.value = header
            cell.font = Font(name='微软雅黑', bold=True)
            cell.fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
            cell.alignment = Alignment(horizontal='center', vertical='center')

        for rule in rules:
            row += 1
            ws[f'A{row}'] = rule.rule_code
            ws[f'B{row}'] = rule.rule_name
            ws[f'C{row}'] = rule.rule_type
            ws[f'D{row}'] = rule.severity

        # 设置列宽
        ws.column_dimensions['A'].width = 20
        ws.column_dimensions['B'].width = 30
        ws.column_dimensions['C'].width = 15
        ws.column_dimensions['D'].width = 15

    @staticmethod
    def _create_error_sheet(ws, errors: List[Dict[str, Any]]):
        """创建错误详情表"""
        # 样式定义
        header_font = Font(name='微软雅黑', bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
        header_alignment = Alignment(horizontal='center', vertical='center')

        # 添加表头
        headers = ['行号', '字段名', '错误值', '错误描述', '严重级别']
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col)
            cell.value = header
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment

        # 添加错误数据
        for row, error in enumerate(errors, start=2):
            ws[f'A{row}'] = error.get('row_number', '')
            ws[f'B{row}'] = error.get('column_name') or error.get('field_name', '')
            ws[f'C{row}'] = str(error.get('original_value', ''))
            ws[f'D{row}'] = error.get('error_message', '')
            ws[f'E{row}'] = error.get('severity', '')

            # 根据严重级别设置颜色
            severity = error.get('severity', 'info')
            if severity == 'error':
                fill_color = 'FFC7CE'  # 红色
            elif severity == 'warning':
                fill_color = 'FFEB9C'  # 黄色
            else:
                fill_color = 'C6EFCE'  # 绿色

            for col in range(1, 6):
                cell = ws.cell(row=row, column=col)
                cell.fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type='solid')
                cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)

        # 如果没有错误数据，添加提示
        if not errors:
            ws['A2'] = '暂无错误数据'
            ws.merge_cells('A2:E2')
            ws['A2'].alignment = Alignment(horizontal='center', vertical='center')
            ws['A2'].font = Font(name='微软雅黑', size=11, italic=True)

        # 设置列宽
        ws.column_dimensions['A'].width = 10
        ws.column_dimensions['B'].width = 20
        ws.column_dimensions['C'].width = 30
        ws.column_dimensions['D'].width = 50
        ws.column_dimensions['E'].width = 15

        # 设置行高
        for row in range(2, len(errors) + 2):
            ws.row_dimensions[row].height = 30

    @staticmethod
    def _create_gdpr_sheet(ws, gdpr_detection_results: List, gdpr_report: Dict[str, Any], task_id: int = None):
        """创建GDPR检测表"""
        # 样式定义
        title_font = Font(name='微软雅黑', size=14, bold=True)
        header_font = Font(name='微软雅黑', bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
        header_alignment = Alignment(horizontal='center', vertical='center')

        # 添加标题
        ws['A1'] = 'GDPR智能检测报告'
        ws['A1'].font = title_font
        ws.merge_cells('A1:F1')
        ws['A1'].alignment = Alignment(horizontal='center', vertical='center')

        # 添加任务信息和审计时间
        row = 2
        if task_id:
            ws[f'A{row}'] = f'任务ID: {task_id}'
            ws[f'A{row}'].font = Font(name='微软雅黑', size=10, italic=True)
        ws[f'D{row}'] = f'审计时间: {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}'
        ws[f'D{row}'].font = Font(name='微软雅黑', size=10, italic=True)

        # 添加合规摘要
        row = 4
        if gdpr_report:
            summary = gdpr_report.get('summary', {})
            ws[f'A{row}'] = '合规分数:'
            ws[f'A{row}'].font = Font(name='微软雅黑', bold=True)
            score = gdpr_report.get('compliance_score', 0)
            ws[f'B{row}'] = f"{score:.1f}/100"
            if score >= 80:
                ws[f'B{row}'].font = Font(name='微软雅黑', bold=True, color='00B050')
            elif score >= 60:
                ws[f'B{row}'].font = Font(name='微软雅黑', bold=True, color='FFC000')
            else:
                ws[f'B{row}'].font = Font(name='微软雅黑', bold=True, color='FF0000')
            row += 1

            ws[f'A{row}'] = '检测项总数:'
            ws[f'A{row}'].font = Font(name='微软雅黑', bold=True)
            ws[f'B{row}'] = summary.get('total_detections', 0)
            row += 1

            ws[f'A{row}'] = '唯一数据类型:'
            ws[f'A{row}'].font = Font(name='微软雅黑', bold=True)
            ws[f'B{row}'] = summary.get('unique_types', 0)
            row += 1

            # 风险分布
            risk_breakdown = summary.get('detection_by_risk', {})
            ws[f'A{row}'] = '风险分布:'
            ws[f'A{row}'].font = Font(name='微软雅黑', bold=True)
            row += 1
            for risk_level, count in risk_breakdown.items():
                ws[f'B{row}'] = f"  {risk_level}: {count}项"
                row += 1

        # 添加检测详情表头
        row += 1
        ws[f'A{row}'] = 'GDPR检测详情'
        ws[f'A{row}'].font = header_font
        ws[f'A{row}'].fill = header_fill
        ws[f'A{row}'].alignment = header_alignment
        ws.merge_cells(f'A{row}:F{row}')

        row += 1
        headers = ['数据类型', '匹配值', '位置', '风险等级', '置信度', '建议']
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=row, column=col)
            cell.value = header
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment

        # 添加检测数据
        for detection in gdpr_detection_results:
            row += 1
            ws[f'A{row}'] = detection.detection_type.value
            ws[f'B{row}'] = detection.matched_value[:50] if detection.matched_value else ''  # 限制长度
            ws[f'C{row}'] = detection.location[:30] if detection.location else ''
            ws[f'D{row}'] = detection.risk_level.value
            ws[f'E{row}'] = f"{detection.confidence:.2f}"
            ws[f'F{row}'] = detection.recommendation[:100] if detection.recommendation else ''

            # 根据风险等级设置颜色
            risk_level = detection.risk_level.value
            if risk_level == '严重':
                fill_color = 'FFC7CE'  # 红色
            elif risk_level == '高':
                fill_color = 'FFD9B3'  # 橙色
            elif risk_level == '中':
                fill_color = 'FFEB9C'  # 黄色
            else:
                fill_color = 'C6EFCE'  # 绿色

            for col in range(1, 7):
                cell = ws.cell(row=row, column=col)
                cell.fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type='solid')
                cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)

        # 设置列宽
        ws.column_dimensions['A'].width = 20
        ws.column_dimensions['B'].width = 30
        ws.column_dimensions['C'].width = 25
        ws.column_dimensions['D'].width = 12
        ws.column_dimensions['E'].width = 10
        ws.column_dimensions['F'].width = 40

        # 设置行高
        for r in range(row - len(gdpr_detection_results) + 1, row + 1):
            ws.row_dimensions[r].height = 30

    @staticmethod
    def _create_rule_evidence_sheet(ws, rule_statistics: List[Dict[str, Any]]):
        """???????"""
        title_font = Font(name='??????', size=14, bold=True)
        header_font = Font(name='??????', bold=True)
        row = 1

        if not rule_statistics:
            ws['A1'] = '????????'
            ws['A1'].font = title_font
            return

        for stat in rule_statistics:
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
            ws[f'A{row}'] = f"{stat.get('rule_code')} - {stat.get('rule_name')}"
            ws[f'A{row}'].font = title_font
            row += 2

            metrics = stat.get('metrics') or {}
            for label, value in metrics.items():
                ws[f'A{row}'] = label
                ws[f'B{row}'] = value
                ws[f'A{row}'].font = header_font
                row += 1

            tables = stat.get('tables') or []
            for table in tables:
                headers = table.get('headers') or []
                rows = table.get('rows') or []
                row = ReportGenerator._write_table(ws, row + 1, table.get('title', ''), headers, rows)

            samples = stat.get('samples') or []
            if samples:
                sample_headers = ['??', '??', '?', '??', '??ID', '??ID', '??ID']
                sample_rows = []
                for sample in samples[:20]:
                    sample_rows.append([
                        sample.get('row_number'),
                        sample.get('field'),
                        sample.get('value'),
                        sample.get('message'),
                        sample.get('record_id'),
                        sample.get('subject_id'),
                        sample.get('dsar_request_id')
                    ])
                row = ReportGenerator._write_table(ws, row + 1, '????', sample_headers, sample_rows)

            row += 2

    @staticmethod
    def _write_table(ws, start_row: int, title: str, headers: List[str], rows: List[List[Any]]) -> int:
        """??????"""
        if title:
            ws[f'A{start_row}'] = title
            ws[f'A{start_row}'].font = Font(name='??????', bold=True)
            start_row += 1

        header_font = Font(name='??????', bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='5B9BD5', end_color='5B9BD5', fill_type='solid')
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=start_row, column=col)
            cell.value = header
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center', vertical='center')

        for data_row in rows:
            start_row += 1
            for col, value in enumerate(data_row, start=1):
                cell = ws.cell(row=start_row, column=col)
                cell.value = value
                cell.alignment = Alignment(horizontal='left', vertical='center')

        return start_row + 2

