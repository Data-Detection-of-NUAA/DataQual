"""
文件解析器
支持多种文件格式的解析：txt, csv, xlsx, json, xml, docx, pdf, doc, zip, rar, 7z等
智能识别GDPR审计相关内容
"""
import csv
import json
import xml.etree.ElementTree as ET
from typing import Any, List, Dict, Optional, Union
import pandas as pd
import os
import tempfile
import shutil
from pathlib import Path


class FileParser:
    """文件解析器类 - 支持多种格式和压缩文件"""

    # GDPR相关关键词
    GDPR_KEYWORDS = [
        'gdpr', 'personal data', 'data subject', 'consent', 'processing',
        'controller', 'processor', 'legitimate interest', 'legal basis',
        '个人数据', '数据主体', '同意', '处理', '控制者', '合法基础',
        'email', 'phone', 'address', 'name', 'id', 'privacy',
        '邮箱', '手机', '地址', '姓名', '身份证', '隐私'
    ]

    @staticmethod
    def parse_file(file_path: str, file_type: str = None) -> Any:
        """
        根据文件类型解析文件，自动识别GDPR相关内容

        Args:
            file_path: 文件路径
            file_type: 文件类型 (可选，自动检测)

        Returns:
            解析后的数据
        """
        # 自动检测文件类型
        if not file_type:
            file_type = Path(file_path).suffix.lower().replace('.', '')
        else:
            file_type = file_type.lower().replace('.', '')

        print(f"解析文件: {file_path}, 类型: {file_type}")

        # 压缩文件：先解压再解析
        if file_type in ['zip', 'rar', '7z', 'tar', 'gz']:
            return FileParser._parse_archive(file_path, file_type)

        # 文本类文件
        elif file_type == 'txt':
            return FileParser._parse_txt(file_path)

        # 结构化数据文件
        elif file_type == 'csv':
            return FileParser._parse_csv(file_path)
        elif file_type in ['xlsx', 'xls']:
            return FileParser._parse_xlsx(file_path)
        elif file_type == 'json':
            return FileParser._parse_json(file_path)
        elif file_type == 'xml':
            return FileParser._parse_xml(file_path)

        # 文档文件
        elif file_type == 'docx':
            return FileParser._parse_docx(file_path)
        elif file_type == 'doc':
            return FileParser._parse_doc(file_path)
        elif file_type == 'pdf':
            return FileParser._parse_pdf(file_path)

        # RTF、HTML等其他格式
        elif file_type == 'rtf':
            return FileParser._parse_rtf(file_path)
        elif file_type in ['html', 'htm']:
            return FileParser._parse_html(file_path)

        else:
            # 尝试作为文本文件读取
            try:
                return FileParser._parse_txt(file_path)
            except Exception as e:
                raise ValueError(f"不支持的文件类型: {file_type}, 错误: {str(e)}")

    @staticmethod
    def _parse_archive(file_path: str, archive_type: str) -> Dict[str, Any]:
        """
        解析压缩文件，提取所有文件并解析

        Args:
            file_path: 压缩文件路径
            archive_type: 压缩类型 (zip/rar/7z/tar/gz)

        Returns:
            包含所有解析结果的字典
        """
        temp_dir = tempfile.mkdtemp()
        results = {
            "archive_type": archive_type,
            "archive_path": file_path,
            "extracted_files": [],
            "parsed_content": {},
            "gdpr_relevant_files": []
        }

        try:
            # 解压文件
            if archive_type == 'zip':
                import zipfile
                with zipfile.ZipFile(file_path, 'r') as zip_ref:
                    zip_ref.extractall(temp_dir)
                    results["extracted_files"] = zip_ref.namelist()

            elif archive_type == 'rar':
                try:
                    import rarfile
                    with rarfile.RarFile(file_path, 'r') as rar_ref:
                        rar_ref.extractall(temp_dir)
                        results["extracted_files"] = rar_ref.namelist()
                except ImportError:
                    raise ImportError("请安装 rarfile: pip install rarfile")

            elif archive_type == '7z':
                try:
                    import py7zr
                    with py7zr.SevenZipFile(file_path, 'r') as archive:
                        archive.extractall(temp_dir)
                        results["extracted_files"] = archive.getnames()
                except ImportError:
                    raise ImportError("请安装 py7zr: pip install py7zr")

            elif archive_type in ['tar', 'gz']:
                import tarfile
                with tarfile.open(file_path, 'r:*') as tar_ref:
                    tar_ref.extractall(temp_dir)
                    results["extracted_files"] = tar_ref.getnames()

            # 遍历解压后的文件
            for root, dirs, files in os.walk(temp_dir):
                for filename in files:
                    file_full_path = os.path.join(root, filename)
                    relative_path = os.path.relpath(file_full_path, temp_dir)

                    try:
                        # 递归解析每个文件
                        parsed = FileParser.parse_file(file_full_path)
                        results["parsed_content"][relative_path] = parsed

                        # 检查是否为GDPR相关文件
                        if FileParser._is_gdpr_relevant(parsed, filename):
                            results["gdpr_relevant_files"].append(relative_path)

                    except Exception as e:
                        print(f"解析文件 {relative_path} 失败: {str(e)}")
                        results["parsed_content"][relative_path] = {
                            "error": str(e),
                            "status": "failed"
                        }

            return results

        finally:
            # 清理临时目录
            shutil.rmtree(temp_dir, ignore_errors=True)

    @staticmethod
    def _is_gdpr_relevant(content: Any, filename: str = "") -> bool:
        """
        判断内容是否与GDPR相关

        Args:
            content: 文件内容
            filename: 文件名

        Returns:
            是否相关
        """
        # 检查文件名
        filename_lower = filename.lower()
        if any(keyword in filename_lower for keyword in ['gdpr', 'privacy', 'consent', 'personal', '隐私', '个人']):
            return True

        # 检查内容
        if isinstance(content, str):
            content_lower = content.lower()
            # 统计关键词出现次数
            keyword_count = sum(1 for keyword in FileParser.GDPR_KEYWORDS if keyword.lower() in content_lower)
            return keyword_count >= 3  # 至少出现3个GDPR关键词

        elif isinstance(content, list):
            # 对于列表数据（如CSV），检查字段名
            if content and isinstance(content[0], dict):
                field_names = ' '.join(content[0].keys()).lower()
                keyword_count = sum(1 for keyword in FileParser.GDPR_KEYWORDS if keyword.lower() in field_names)
                return keyword_count >= 2

        elif isinstance(content, dict):
            # 递归检查字典
            dict_str = str(content).lower()
            keyword_count = sum(1 for keyword in FileParser.GDPR_KEYWORDS if keyword.lower() in dict_str)
            return keyword_count >= 3

        return False

    @staticmethod
    def _parse_txt(file_path: str) -> str:
        """解析TXT文件"""
        encodings = ['utf-8', 'gbk', 'gb2312', 'latin1', 'cp1252']

        for encoding in encodings:
            try:
                with open(file_path, 'r', encoding=encoding) as f:
                    return f.read()
            except UnicodeDecodeError:
                continue

        # 如果所有编码都失败，使用二进制模式读取
        with open(file_path, 'rb') as f:
            return f.read().decode('utf-8', errors='ignore')

    @staticmethod
    def _parse_csv(file_path: str) -> List[Dict[str, Any]]:
        """解析CSV文件"""
        encodings = ['utf-8', 'gbk', 'gb2312', 'latin1']

        for encoding in encodings:
            try:
                df = pd.read_csv(file_path, encoding=encoding)
                # 替换 NaN 为 None
                df = df.where(pd.notna(df), None)
                return df.to_dict('records')
            except (UnicodeDecodeError, pd.errors.ParserError):
                continue

        raise ValueError(f"无法解析CSV文件: {file_path}")

    @staticmethod
    def _parse_xlsx(file_path: str) -> List[Dict[str, Any]]:
        """解析Excel文件"""
        df = pd.read_excel(file_path)
        # 替换 NaN 为 None
        df = df.where(pd.notna(df), None)
        return df.to_dict('records')

    @staticmethod
    def _parse_json(file_path: str) -> Any:
        """解析JSON文件"""
        encodings = ['utf-8', 'gbk', 'gb2312']

        for encoding in encodings:
            try:
                with open(file_path, 'r', encoding=encoding) as f:
                    return json.load(f)
            except (UnicodeDecodeError, json.JSONDecodeError):
                continue

        raise ValueError(f"无法解析JSON文件: {file_path}")

    @staticmethod
    def _parse_xml(file_path: str) -> Dict[str, Any]:
        """解析XML文件，返回结构化数据"""
        tree = ET.parse(file_path)
        root = tree.getroot()

        # 转换为字典结构
        def xml_to_dict(element):
            result = {}

            # 添加属性
            if element.attrib:
                result['@attributes'] = element.attrib

            # 添加文本内容
            if element.text and element.text.strip():
                result['#text'] = element.text.strip()

            # 添加子元素
            for child in element:
                child_data = xml_to_dict(child)
                if child.tag in result:
                    # 如果标签已存在，转换为列表
                    if not isinstance(result[child.tag], list):
                        result[child.tag] = [result[child.tag]]
                    result[child.tag].append(child_data)
                else:
                    result[child.tag] = child_data

            return result

        return {root.tag: xml_to_dict(root)}

    @staticmethod
    def _parse_docx(file_path: str) -> str:
        """解析DOCX文件"""
        try:
            from docx import Document
            doc = Document(file_path)

            # 提取段落文本
            paragraphs = [para.text for para in doc.paragraphs if para.text.strip()]

            # 提取表格数据
            tables_text = []
            for table in doc.tables:
                for row in table.rows:
                    row_text = '\t'.join([cell.text for cell in row.cells])
                    if row_text.strip():
                        tables_text.append(row_text)

            # 合并所有内容
            all_text = '\n'.join(paragraphs)
            if tables_text:
                all_text += '\n\n=== Tables ===\n' + '\n'.join(tables_text)

            return all_text

        except ImportError:
            raise ImportError("请安装 python-docx: pip install python-docx")

    @staticmethod
    def _parse_doc(file_path: str) -> str:
        """解析DOC文件（旧版Word）"""
        try:
            import textract
            # textract可以处理各种文档格式
            text = textract.process(file_path).decode('utf-8', errors='ignore')
            return text
        except ImportError:
            # 降级方案：使用antiword（Linux）或尝试转换
            try:
                import subprocess
                result = subprocess.run(['antiword', file_path],
                                      capture_output=True,
                                      text=True,
                                      timeout=30)
                if result.returncode == 0:
                    return result.stdout
            except Exception:
                pass

            raise ImportError(
                "请安装 textract: pip install textract\n"
                "或在Linux上安装 antiword: sudo apt-get install antiword"
            )

    @staticmethod
    def _parse_pdf(file_path: str) -> str:
        """解析PDF文件"""
        try:
            import pdfplumber
            text_content = []

            with pdfplumber.open(file_path) as pdf:
                for page_num, page in enumerate(pdf.pages, 1):
                    # 提取文本
                    page_text = page.extract_text()
                    if page_text:
                        text_content.append(f"=== Page {page_num} ===\n{page_text}")

                    # 提取表格
                    tables = page.extract_tables()
                    for table_num, table in enumerate(tables, 1):
                        if table:
                            table_text = '\n'.join(['\t'.join([str(cell) if cell else '' for cell in row]) for row in table])
                            text_content.append(f"\n[Table {table_num} on Page {page_num}]\n{table_text}")

            return '\n\n'.join(text_content)

        except ImportError:
            raise ImportError("请安装 pdfplumber: pip install pdfplumber")

    @staticmethod
    def _parse_rtf(file_path: str) -> str:
        """解析RTF文件"""
        try:
            from striprtf.striprtf import rtf_to_text
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                rtf_content = f.read()
            return rtf_to_text(rtf_content)
        except ImportError:
            raise ImportError("请安装 striprtf: pip install striprtf")

    @staticmethod
    def _parse_html(file_path: str) -> str:
        """解析HTML文件"""
        try:
            from bs4 import BeautifulSoup

            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                html_content = f.read()

            soup = BeautifulSoup(html_content, 'html.parser')

            # 移除脚本和样式
            for script in soup(["script", "style"]):
                script.decompose()

            # 提取文本
            text = soup.get_text()

            # 清理空白
            lines = (line.strip() for line in text.splitlines())
            chunks = (phrase.strip() for line in lines for phrase in line.split("  "))
            text = '\n'.join(chunk for chunk in chunks if chunk)

            return text

        except ImportError:
            raise ImportError("请安装 beautifulsoup4: pip install beautifulsoup4")

    @staticmethod
    def extract_gdpr_content(content: Any, content_type: str = "text") -> Dict[str, Any]:
        """
        从解析后的内容中提取GDPR相关信息

        Args:
            content: 解析后的内容
            content_type: 内容类型 (text/structured/archive)

        Returns:
            提取的GDPR相关信息
        """
        result = {
            "content_type": content_type,
            "has_gdpr_content": False,
            "detected_keywords": [],
            "detected_data_types": [],
            "confidence_score": 0.0,
            "summary": ""
        }

        if content_type == "text" and isinstance(content, str):
            # 文本内容分析
            content_lower = content.lower()

            # 检测关键词
            detected_keywords = [kw for kw in FileParser.GDPR_KEYWORDS if kw.lower() in content_lower]
            result["detected_keywords"] = detected_keywords

            # 检测数据类型
            data_types = []
            if any(kw in content_lower for kw in ['email', '邮箱', '@']):
                data_types.append('email')
            if any(kw in content_lower for kw in ['phone', 'mobile', '手机', '电话']):
                data_types.append('phone')
            if any(kw in content_lower for kw in ['address', '地址']):
                data_types.append('address')
            if any(kw in content_lower for kw in ['name', '姓名', '名称']):
                data_types.append('name')
            if any(kw in content_lower for kw in ['id', 'idcard', '身份证', '证件']):
                data_types.append('idcard')

            result["detected_data_types"] = data_types

            # 计算置信度
            keyword_score = min(len(detected_keywords) / 10, 1.0) * 0.6
            data_type_score = min(len(data_types) / 5, 1.0) * 0.4
            result["confidence_score"] = keyword_score + data_type_score

            result["has_gdpr_content"] = result["confidence_score"] > 0.3

            # 生成摘要（提取前500字符）
            result["summary"] = content[:500] + "..." if len(content) > 500 else content

        elif content_type == "structured" and isinstance(content, list):
            # 结构化数据分析
            if content and isinstance(content[0], dict):
                field_names = list(content[0].keys())
                field_names_lower = [f.lower() for f in field_names]

                # 检测字段类型
                data_types = []
                if any('email' in f or 'mail' in f for f in field_names_lower):
                    data_types.append('email')
                if any('phone' in f or 'mobile' in f or '手机' in f for f in field_names_lower):
                    data_types.append('phone')
                if any('address' in f or '地址' in f for f in field_names_lower):
                    data_types.append('address')
                if any('name' in f or '姓名' in f for f in field_names_lower):
                    data_types.append('name')
                if any('id' in f or '身份证' in f for f in field_names_lower):
                    data_types.append('idcard')

                result["detected_data_types"] = data_types
                result["confidence_score"] = min(len(data_types) / 5, 1.0)
                result["has_gdpr_content"] = len(data_types) >= 2
                result["summary"] = f"结构化数据，包含 {len(content)} 条记录，{len(field_names)} 个字段: {', '.join(field_names[:10])}"

        elif content_type == "archive" and isinstance(content, dict):
            # 压缩包分析
            gdpr_files = content.get("gdpr_relevant_files", [])
            result["has_gdpr_content"] = len(gdpr_files) > 0
            result["confidence_score"] = min(len(gdpr_files) / 5, 1.0)
            result["summary"] = f"压缩包包含 {len(content.get('extracted_files', []))} 个文件，其中 {len(gdpr_files)} 个与GDPR相关"
            result["detected_keywords"] = gdpr_files

        return result
